"""
案件成员关系与资源范围授权测试。

覆盖：
- 列表/详情/导出/操作日志查询走同一授权判断；
- 只能访问分配给自己的案件与时间段；
- 退出案件后新访问立即停止、历史操作仍可审计；
- 一人多案为并集且互不可见；
- 案件合并/拆分不泄露其他案件记录；
- 拒绝访问落 AccessDecision，管理员可看到所依据的成员关系版本与快照。
"""
from datetime import timedelta
from decimal import Decimal

from django.utils import timezone
from rest_framework.test import APITestCase

from apps.authentication.backends import generate_token
from apps.authentication.models import User, OperationLog
from apps.warehouse.models import (
    Approval, Category, Goods, StockIn, StockOut, Unit, Variety, Warning,
)
from .models import Case, CaseMembership, CaseLineage, AccessDecision
from .services import merge_cases, split_case


class ScopeTestBase(APITestCase):
    def setUp(self):
        self.now = timezone.now()
        self.admin = User.objects.create_user('admin', 'pass123456', role='admin')
        self.alice = User.objects.create_user('alice', 'pass123456', role='user')
        self.bob = User.objects.create_user('bob', 'pass123456', role='user')

        self.case_a = Case.objects.create(name='专案甲', case_no='A-001', created_by=self.admin)
        self.case_b = Case.objects.create(name='专案乙', case_no='B-001', created_by=self.admin)

        unit = Unit.objects.create(name='件', created_by=self.admin)
        category = Category.objects.create(name='管控器材', unit=unit, created_by=self.admin)
        variety = Variety.objects.create(name='记录设备', category=category, created_by=self.admin)
        self.variety = variety

        # 甲案：两条货物，分别落在授权窗口内/外
        self.goods_in = self._goods('G-IN', '窗口内货物', self.case_a, days_ago=3)
        self.goods_before = self._goods('G-BEFORE', '窗口前货物', self.case_a, days_ago=20)
        # 乙案：alice 无权
        self.goods_b = self._goods('G-B', '乙案货物', self.case_b, days_ago=1)

    def _goods(self, code, name, case, days_ago):
        goods = Goods.objects.create(
            case=case, variety=self.variety, name=name, code=code,
            quantity=Decimal('1'), warning_threshold=Decimal('5'),
        )
        Goods.objects.filter(pk=goods.pk).update(created_at=self.now - timedelta(days=days_ago))
        goods.refresh_from_db()
        return goods

    def grant(self, case, user, days_from=None, days_to=None):
        vf = self.now + timedelta(days=days_from) if days_from is not None else self.now - timedelta(days=100)
        vt = None if days_to is None else self.now + timedelta(days=days_to)
        return case.grant_membership(user, valid_from=vf, valid_to=vt, granted_by=self.admin)

    def as_user(self, user):
        self.client.credentials(HTTP_AUTHORIZATION=f'Bearer {generate_token(user)}')


class ListAndDetailScopeTest(ScopeTestBase):
    def test_member_list_only_own_case_within_window(self):
        self.grant(self.case_a, self.alice, days_from=-10, days_to=10)
        self.as_user(self.alice)

        resp = self.client.get('/api/goods/')
        self.assertEqual(resp.status_code, 200)
        codes = {item['code'] for item in resp.json()['data']['list']}
        self.assertEqual(codes, {'G-IN'})  # 不见窗口前货物，也不见乙案货物

    def test_member_detail_allow_and_deny_other_case(self):
        self.grant(self.case_a, self.alice, days_from=-10, days_to=10)
        self.as_user(self.alice)

        ok = self.client.get(f'/api/goods/{self.goods_in.id}/')
        deny_other = self.client.get(f'/api/goods/{self.goods_b.id}/')
        deny_window = self.client.get(f'/api/goods/{self.goods_before.id}/')

        self.assertEqual(ok.status_code, 200)
        self.assertEqual(deny_other.status_code, 403)
        self.assertEqual(deny_window.status_code, 403)

    def test_no_membership_sees_nothing(self):
        self.as_user(self.bob)
        resp = self.client.get('/api/goods/')
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(resp.json()['data']['total'], 0)
        self.assertEqual(self.client.get(f'/api/goods/{self.goods_in.id}/').status_code, 403)

    def test_future_window_is_not_effective(self):
        # 授权从明天开始：今天一切访问都被拒绝
        self.grant(self.case_a, self.alice, days_from=1, days_to=10)
        self.as_user(self.alice)
        self.assertEqual(self.client.get('/api/goods/').json()['data']['total'], 0)
        self.assertEqual(self.client.get(f'/api/goods/{self.goods_in.id}/').status_code, 403)

    def test_admin_sees_all(self):
        self.as_user(self.admin)
        resp = self.client.get('/api/goods/')
        self.assertEqual(resp.json()['data']['total'], 3)

    def test_member_cannot_create_in_other_case(self):
        self.grant(self.case_a, self.alice, days_from=-10, days_to=10)
        self.as_user(self.alice)
        resp = self.client.post('/api/stock-in/', {
            'case': self.case_b.id,
            'goods': self.goods_b.id,
            'quantity': '1',
        }, format='json')
        self.assertEqual(resp.status_code, 403)
        self.assertFalse(StockIn.objects.filter(case=self.case_b, operator=self.alice).exists())

    def test_member_can_create_in_own_case(self):
        self.grant(self.case_a, self.alice, days_from=-10, days_to=10)
        self.as_user(self.alice)
        resp = self.client.post('/api/stock-in/', {
            'case': self.case_a.id,
            'goods': self.goods_in.id,
            'quantity': '2',
        }, format='json')
        self.assertEqual(resp.status_code, 200)
        self.assertTrue(StockIn.objects.filter(case=self.case_a, operator=self.alice).exists())


class MembershipVersioningTest(ScopeTestBase):
    def test_grant_creates_versions_and_keeps_history(self):
        v1 = self.grant(self.case_a, self.alice, days_from=-10, days_to=-2)
        v2 = self.grant(self.case_a, self.alice, days_from=-1, days_to=5)

        self.assertEqual(v1.version, 1)
        self.assertEqual(v2.version, 2)
        v1.refresh_from_db()
        self.assertIsNotNone(v1.valid_to)  # 旧版本被关闭但保留
        self.assertTrue(CaseMembership.objects.filter(pk=v1.pk).exists())

        # 拒绝记录依据最新版本
        self.as_user(self.alice)
        self.client.get(f'/api/goods/{self.goods_before.id}/')  # 20天前，两版窗口都不覆盖
        decision = AccessDecision.objects.filter(
            user=self.alice, decision='deny', resource_type='goods'
        ).latest('created_at')
        self.assertEqual(decision.reason, 'outside_time_window')
        self.assertIsNotNone(decision.membership)
        self.assertEqual(decision.membership.version, 2)
        self.assertEqual(decision.membership_snapshot['version'], 2)
        self.assertIn('valid_from', decision.membership_snapshot)

    def test_exit_stops_new_access_immediately_but_history_auditable(self):
        self.grant(self.case_a, self.alice, days_from=-10)  # 开放窗口
        self.as_user(self.alice)
        # 在案期间登记一条入库（中间件同时写操作日志）
        create = self.client.post('/api/stock-in/', {
            'case': self.case_a.id, 'goods': self.goods_in.id, 'quantity': '1',
        }, format='json')
        self.assertEqual(create.status_code, 200)

        # 管理员将 alice 退出案件
        self.case_a.revoke_membership(self.alice, revoked_by=self.admin)

        # 新的列表、详情访问立即停止
        self.assertEqual(self.client.get('/api/goods/').json()['data']['total'], 0)
        self.assertEqual(self.client.get(f'/api/goods/{self.goods_in.id}/').status_code, 403)

        # 历史操作仍可自查（审计不中断）
        logs = self.client.get('/api/auth/logs/').json()['data']['list']
        self.assertTrue(any(log['user'] == self.alice.id for log in logs))

        # 管理员仍能审计 alice 的历史操作
        self.as_user(self.admin)
        admin_logs = self.client.get('/api/auth/logs/', {'user': self.alice.id}).json()['data']['list']
        self.assertTrue(any(log['user'] == self.alice.id for log in admin_logs))

    def test_revoked_user_denial_records_closed_version(self):
        v = self.grant(self.case_a, self.alice, days_from=-10)
        self.case_a.revoke_membership(self.alice, revoked_by=self.admin)
        self.as_user(self.alice)
        self.client.get(f'/api/goods/{self.goods_in.id}/')
        decision = AccessDecision.objects.filter(user=self.alice, decision='deny').latest('created_at')
        self.assertEqual(decision.reason, 'not_member')
        self.assertEqual(decision.membership_id, v.id)  # 依据的就是被关闭的那一版
        self.assertIsNotNone(decision.membership_snapshot['valid_to'])


class MultiCaseTest(ScopeTestBase):
    def test_user_in_multiple_cases_gets_union_without_leakage(self):
        self.grant(self.case_a, self.alice, days_from=-10, days_to=10)
        self.grant(self.case_b, self.alice, days_from=-10, days_to=10)
        self.as_user(self.alice)

        codes = {item['code'] for item in self.client.get('/api/goods/').json()['data']['list']}
        self.assertEqual(codes, {'G-IN', 'G-B'})  # 两案并集，窗口外的 G-BEFORE 仍不可见

    def test_two_users_do_not_see_each_others_cases(self):
        self.grant(self.case_a, self.alice, days_from=-10, days_to=10)
        self.grant(self.case_b, self.bob, days_from=-10, days_to=10)
        self.as_user(self.alice)
        self.assertEqual(self.client.get(f'/api/goods/{self.goods_b.id}/').status_code, 403)
        self.as_user(self.bob)
        self.assertEqual(self.client.get(f'/api/goods/{self.goods_in.id}/').status_code, 403)


class MergeSplitTest(ScopeTestBase):
    def test_merge_moves_resources_and_closes_access_without_inheriting_members(self):
        self.grant(self.case_a, self.alice, days_from=-10, days_to=10)
        case_c = Case.objects.create(name='承接案', case_no='C-001', created_by=self.admin)

        target, lineages = merge_cases([self.case_a.id], case_c.id, operated_by=self.admin)
        self.assertEqual(target, case_c)
        self.goods_in.refresh_from_db()
        self.assertEqual(self.goods_in.case_id, case_c.id)
        self.case_a.refresh_from_db()
        self.assertEqual(self.case_a.status, Case.STATUS_MERGED)
        self.assertEqual(self.case_a.merged_into_id, case_c.id)
        self.assertEqual(len(lineages), 1)
        self.assertIn(self.goods_in.id, lineages[0].goods_ids)

        # alice 不是承接案成员：即使持有甲案开放版本，访问立即被拒
        self.as_user(self.alice)
        self.assertEqual(self.client.get(f'/api/goods/{self.goods_in.id}/').status_code, 403)
        self.assertEqual(self.client.get('/api/goods/').json()['data']['total'], 0)

        decision = AccessDecision.objects.filter(
            user=self.alice, decision='deny', resource_id=str(self.goods_in.id)
        ).latest('created_at')
        self.assertEqual(decision.reason, 'not_member')
        self.assertEqual(decision.case_id, case_c.id)  # 资源已在承接案，需重新授权

        # 重新授权承接案后恢复
        self.grant(case_c, self.alice, days_from=-10, days_to=10)
        self.assertEqual(self.client.get(f'/api/goods/{self.goods_in.id}/').status_code, 200)

    def test_closed_case_denies_with_case_closed_reason(self):
        self.grant(self.case_a, self.alice, days_from=-10, days_to=10)
        # 案件直接归档、资源留在原案
        self.case_a.status = Case.STATUS_ARCHIVED
        self.case_a.save()
        self.as_user(self.alice)
        self.assertEqual(self.client.get(f'/api/goods/{self.goods_in.id}/').status_code, 403)
        decision = AccessDecision.objects.filter(user=self.alice, decision='deny').latest('created_at')
        self.assertEqual(decision.reason, 'case_closed')
        self.assertIsNotNone(decision.membership)  # 即使成员版本还在，案件关闭也拒绝

    def test_merge_rejects_closed_source_and_self_target(self):
        self.case_a.status = Case.STATUS_ARCHIVED
        self.case_a.save()
        with self.assertRaises(Exception):
            merge_cases([self.case_a.id], self.case_b.id, operated_by=self.admin)
        with self.assertRaises(Exception):
            merge_cases([self.case_b.id], self.case_b.id, operated_by=self.admin)

    def test_split_partitions_resources_without_leakage(self):
        self.grant(self.case_a, self.alice, days_from=-10, days_to=10)
        source, created, lineages = split_case(
            self.case_a.id,
            [{
                'name': '拆分子案', 'case_no': 'A-S1',
                'goods_ids': [self.goods_in.id],
                'stock_in_ids': [], 'stock_out_ids': [],
            }],
            operated_by=self.admin,
        )
        self.goods_in.refresh_from_db()
        self.goods_before.refresh_from_db()
        self.assertEqual(self.goods_in.case_id, created[0].id)
        self.assertEqual(self.goods_before.case_id, self.case_a.id)  # 未分配的留在原案
        self.assertEqual(lineages[0].relation, 'split')

        # 子案未授权给 alice：看不到被拆走的货物
        self.as_user(self.alice)
        self.assertEqual(self.client.get(f'/api/goods/{self.goods_in.id}/').status_code, 403)
        # 留在原案、窗口外的货物依旧不可见；原案窗口内已无资源
        self.assertEqual(self.client.get('/api/goods/').json()['data']['total'], 0)

    def test_split_rejects_foreign_resource(self):
        with self.assertRaises(Exception):
            split_case(
                self.case_a.id,
                [{
                    'name': '越权子案', 'case_no': 'A-S2',
                    'goods_ids': [self.goods_b.id],
                }],
                operated_by=self.admin,
            )
        self.goods_b.refresh_from_db()
        self.assertEqual(self.goods_b.case_id, self.case_b.id)  # 整体回滚


class ExportAndLogScopeTest(ScopeTestBase):
    def test_export_is_scoped_by_same_authorizer(self):
        import io
        from openpyxl import load_workbook

        self.grant(self.case_a, self.alice, days_from=-10, days_to=10)
        self.as_user(self.alice)
        resp = self.client.get('/api/export/', {'type': 'goods'})
        self.assertEqual(resp.status_code, 200)

        wb = load_workbook(io.BytesIO(resp.content))
        ws = wb.active
        codes = {row[0] for row in ws.iter_rows(min_row=2, values_only=True)}
        self.assertIn('G-IN', codes)
        self.assertNotIn('G-B', codes)
        self.assertNotIn('G-BEFORE', codes)
        decision = AccessDecision.objects.filter(user=self.alice, action='export').first()
        self.assertIsNotNone(decision)

    def test_operation_log_query_is_scoped(self):
        self.grant(self.case_a, self.alice, days_from=-10, days_to=10)
        self.grant(self.case_b, self.bob, days_from=-10, days_to=10)
        # bob 在乙案产生一条带案件上下文的操作
        self.as_user(self.bob)
        self.client.post('/api/stock-in/', {
            'case': self.case_b.id, 'goods': self.goods_b.id, 'quantity': '1',
        }, format='json')

        # alice 查日志：看不到 bob 在乙案的操作
        self.as_user(self.alice)
        logs = self.client.get('/api/auth/logs/').json()['data']['list']
        self.assertFalse(any(log['user'] == self.bob.id for log in logs))

        # 管理员看得到全部
        self.as_user(self.admin)
        admin_logs = self.client.get('/api/auth/logs/').json()['data']['list']
        self.assertTrue(any(log['user'] == self.bob.id for log in admin_logs))


class IndirectOwnershipScopeTest(ScopeTestBase):
    """预警（经货物）、审批（经出库记录）沿用同一授权判断"""

    def setUp(self):
        super().setUp()
        self.stock_out_b = StockOut.objects.create(
            case=self.case_b, goods=self.goods_b, operator=self.bob,
            receiver='领用人', quantity=Decimal('1'),
        )
        self.warning_in = Warning.objects.create(
            goods=self.goods_in, type='low_stock', message='甲案预警'
        )
        self.warning_b = Warning.objects.create(
            goods=self.goods_b, type='low_stock', message='乙案预警'
        )
        self.approval_b = Approval.objects.create(
            stock_out=self.stock_out_b, approver=self.admin,
        )

    def test_warning_scoped_via_goods_case(self):
        self.grant(self.case_a, self.alice, days_from=-10, days_to=10)
        self.as_user(self.alice)

        rows = self.client.get('/api/warnings/').json()['data']['list']
        self.assertEqual({r['id'] for r in rows}, {self.warning_in.id})

        self.assertEqual(self.client.get(f'/api/warnings/{self.warning_b.id}/').status_code, 403)
        self.assertEqual(self.client.get(f'/api/warnings/{self.warning_in.id}/').status_code, 200)

    def test_approval_scoped_via_stock_out_case(self):
        self.grant(self.case_b, self.bob, days_from=-10, days_to=10)
        self.as_user(self.bob)

        rows = self.client.get('/api/approvals/').json()['data']['list']
        self.assertEqual({r['id'] for r in rows}, {self.approval_b.id})
        self.assertEqual(self.client.get(f'/api/approvals/{self.approval_b.id}/').status_code, 200)

        self.as_user(self.alice)
        self.assertEqual(self.client.get(f'/api/approvals/{self.approval_b.id}/').status_code, 403)


class AccessDecisionAdminTest(ScopeTestBase):
    def test_only_admin_can_list_decisions(self):
        self.as_user(self.alice)
        self.assertEqual(self.client.get('/api/access-decisions/').status_code, 403)

        # 产生一次拒绝
        self.client.get(f'/api/goods/{self.goods_b.id}/')

        self.as_user(self.admin)
        resp = self.client.get('/api/access-decisions/', {'decision': 'deny'})
        self.assertEqual(resp.status_code, 200)
        denied = resp.json()['data']['list']
        self.assertTrue(any(d['resource_type'] == 'goods' for d in denied))

    def test_non_admin_cannot_grant_membership(self):
        self.as_user(self.alice)
        resp = self.client.post(
            f'/api/cases/{self.case_a.id}/memberships/',
            {'user_id': self.bob.id}, format='json',
        )
        self.assertEqual(resp.status_code, 403)

    def test_case_list_scoped_for_member(self):
        self.grant(self.case_a, self.alice, days_from=-10, days_to=10)
        self.as_user(self.alice)
        case_nos = {c['case_no'] for c in self.client.get('/api/cases/').json()['data']['list']}
        self.assertEqual(case_nos, {'A-001'})
