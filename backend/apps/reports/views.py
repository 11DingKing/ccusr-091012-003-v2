"""
报表视图
"""
import logging
import os
import shutil
from datetime import datetime, timedelta
from django.db.models import Sum, Count, Q, F
from django.http import HttpResponse
from rest_framework.views import APIView
from rest_framework.permissions import IsAuthenticated
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from apps.core.response import success_response, error_response
from apps.warehouse.models import Goods, StockIn, StockOut, Warning
from .models import DailyReport

logger = logging.getLogger('apps')


class DashboardView(APIView):
    """仪表盘数据（按案件成员关系范围聚合，管理员为全量）"""
    permission_classes = [IsAuthenticated]

    def get(self, request):
        from apps.cases import http as case_http

        today = datetime.now().date()

        goods_qs = case_http.visible_queryset(
            request, Goods.objects.all(), 'goods', action='dashboard')
        stock_in_qs = case_http.visible_queryset(
            request, StockIn.objects.all(), 'stock_in', action='dashboard')
        stock_out_qs = case_http.visible_queryset(
            request, StockOut.objects.all(), 'stock_out', action='dashboard')
        warning_qs = case_http.visible_queryset(
            request, Warning.objects.all(), 'warning', action='dashboard')

        # 货物统计
        goods_count = goods_qs.filter(is_active=True).count()
        warning_goods_count = goods_qs.filter(
            is_active=True,
            quantity__lte=F('warning_threshold')
        ).count()

        # 今日入库统计
        today_in = stock_in_qs.filter(
            stock_in_time__date=today
        ).aggregate(
            count=Count('id'),
            total=Sum('quantity')
        )

        # 今日出库统计
        today_out = stock_out_qs.filter(
            stock_out_time__date=today,
            status='completed'
        ).aggregate(
            count=Count('id'),
            total=Sum('quantity')
        )

        # 待审批数量
        pending_approval_count = case_http.visible_queryset(
            request, StockOut.objects.all(), 'stock_out', action='dashboard'
        ).filter(status='pending').count()

        # 未读预警数量
        unread_warning_count = warning_qs.filter(is_read=False).count()

        # 最近7天入库趋势
        in_trend = []
        out_trend = []
        for i in range(6, -1, -1):
            date = today - timedelta(days=i)
            in_data = stock_in_qs.filter(
                stock_in_time__date=date
            ).aggregate(total=Sum('quantity'))
            out_data = stock_out_qs.filter(
                stock_out_time__date=date,
                status='completed'
            ).aggregate(total=Sum('quantity'))

            in_trend.append({
                'date': date.strftime('%m-%d'),
                'value': float(in_data['total'] or 0)
            })
            out_trend.append({
                'date': date.strftime('%m-%d'),
                'value': float(out_data['total'] or 0)
            })

        # 库存预警货物列表
        warning_goods = goods_qs.filter(
            is_active=True,
            quantity__lte=F('warning_threshold')
        ).values('id', 'name', 'code', 'quantity', 'warning_threshold')[:5]

        return success_response(data={
            'goods_count': goods_count,
            'warning_goods_count': warning_goods_count,
            'today_in_count': today_in['count'] or 0,
            'today_in_total': float(today_in['total'] or 0),
            'today_out_count': today_out['count'] or 0,
            'today_out_total': float(today_out['total'] or 0),
            'pending_approval_count': pending_approval_count,
            'unread_warning_count': unread_warning_count,
            'in_trend': in_trend,
            'out_trend': out_trend,
            'warning_goods': list(warning_goods)
        })


class DailyReportView(APIView):
    """每日报表（按案件成员关系范围统计）"""
    permission_classes = [IsAuthenticated]

    def get(self, request):
        from apps.cases import http as case_http

        scope, case_id, denied = case_http.guard_case_param(
            request, 'stock_in', action='report')
        if denied:
            return denied

        start_date = request.query_params.get('start_date')
        end_date = request.query_params.get('end_date')

        if not start_date:
            start_date = (datetime.now() - timedelta(days=30)).date()
        else:
            start_date = datetime.strptime(start_date, '%Y-%m-%d').date()

        if not end_date:
            end_date = datetime.now().date()
        else:
            end_date = datetime.strptime(end_date, '%Y-%m-%d').date()

        stock_in_qs = case_http.visible_queryset(
            request, StockIn.objects.all(), 'stock_in', action='report')
        stock_out_qs = case_http.visible_queryset(
            request, StockOut.objects.all(), 'stock_out', action='report')
        warning_qs = case_http.visible_queryset(
            request, Warning.objects.all(), 'warning', action='report')
        if case_id:
            stock_in_qs = stock_in_qs.filter(goods__case_id=case_id)
            stock_out_qs = stock_out_qs.filter(goods__case_id=case_id)
            warning_qs = warning_qs.filter(goods__case_id=case_id)

        reports = []
        current_date = start_date

        while current_date <= end_date:
            # 入库统计
            in_data = stock_in_qs.filter(
                stock_in_time__date=current_date
            ).aggregate(
                count=Count('id'),
                total=Sum('quantity')
            )

            # 出库统计
            out_data = stock_out_qs.filter(
                stock_out_time__date=current_date,
                status='completed'
            ).aggregate(
                count=Count('id'),
                total=Sum('quantity')
            )

            # 预警统计
            warning_count = warning_qs.filter(
                created_at__date=current_date
            ).count()

            reports.append({
                'date': current_date.strftime('%Y-%m-%d'),
                'in_count': in_data['count'] or 0,
                'in_total': float(in_data['total'] or 0),
                'out_count': out_data['count'] or 0,
                'out_total': float(out_data['total'] or 0),
                'warning_count': warning_count
            })

            current_date += timedelta(days=1)

        return success_response(data=reports)


class ExportView(APIView):
    """
    数据导出（与列表、详情使用同一授权判断）

    普通成员只能导出其当前授权案件范围内的记录，可显式传 ?case= 限定单个案件；
    管理员可导出全部。
    """
    permission_classes = [IsAuthenticated]

    def get(self, request):
        from apps.cases import http as case_http

        export_type = request.query_params.get('type', 'goods')

        if export_type == 'goods':
            _, case_id, denied = case_http.guard_case_param(request, 'goods', action='export')
            if denied:
                return denied
            qs = case_http.visible_queryset(
                request,
                Goods.objects.select_related(
                    'case', 'variety', 'variety__category', 'variety__category__unit'),
                'goods', action='export')
            if case_id:
                qs = qs.filter(case_id=case_id)
            goods_list = qs.filter(is_active=True)
            return self._export_goods(goods_list)

        if export_type == 'stock_in':
            _, case_id, denied = case_http.guard_case_param(request, 'stock_in', action='export')
            if denied:
                return denied
            qs = case_http.visible_queryset(
                request,
                StockIn.objects.select_related(
                    'goods', 'goods__case', 'goods__variety__category__unit', 'operator'),
                'stock_in', action='export')
            if case_id:
                qs = qs.filter(goods__case_id=case_id)
            return self._export_stock_in(qs)

        if export_type == 'stock_out':
            _, case_id, denied = case_http.guard_case_param(request, 'stock_out', action='export')
            if denied:
                return denied
            qs = case_http.visible_queryset(
                request,
                StockOut.objects.select_related(
                    'goods', 'goods__case', 'goods__variety__category__unit', 'operator'),
                'stock_out', action='export')
            if case_id:
                qs = qs.filter(goods__case_id=case_id)
            return self._export_stock_out(qs)

        return error_response(message='无效的导出类型')

    # ---- 各类型导出的拼装（数据已在授权范围内） ----

    @staticmethod
    def _styles(wb):
        ws = wb.active
        header_font = Font(bold=True, color='FFFFFF')
        header_fill = PatternFill(start_color='0066FF', end_color='0066FF', fill_type='solid')
        header_alignment = Alignment(horizontal='center', vertical='center')
        thin_border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
        return ws, header_font, header_fill, header_alignment, thin_border

    @staticmethod
    def _autosize(ws):
        for column in ws.columns:
            max_length = 0
            column_letter = column[0].column_letter
            for cell in column:
                try:
                    if cell.value is not None and len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except Exception:
                    pass
            ws.column_dimensions[column_letter].width = (max_length + 2) * 1.2

    def _finalize(self, ws, headers, border, font, fill, alignment, filename):
        for col in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col)
            cell.font = font
            cell.fill = fill
            cell.alignment = alignment
            cell.border = border
        self._autosize(ws)
        response = HttpResponse(
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = f'attachment; filename="{filename}"'
        wb = ws.parent
        wb.save(response)
        logger.info("Data exported: %s rows=%s", filename, ws.max_row - 1)
        return response

    def _export_goods(self, goods_list):
        wb = Workbook()
        ws, font, fill, alignment, border = self._styles(wb)
        ws.title = '货物清单'
        headers = ['案件编号', '编码', '名称', '品类', '品种', '规格',
                   '库存', '单位', '存放位置', '预警阈值']
        ws.append(headers)
        for goods in goods_list:
            unit_name = (goods.variety.category.unit.name
                         if goods.variety and goods.variety.category
                         and goods.variety.category.unit else '')
            ws.append([
                goods.case.case_no if goods.case else '',
                goods.code,
                goods.name,
                goods.variety.category.name if goods.variety else '',
                goods.variety.name if goods.variety else '',
                goods.specification,
                float(goods.quantity),
                unit_name,
                goods.location,
                float(goods.warning_threshold)
            ])
        filename = f'货物清单_{datetime.now().strftime("%Y%m%d%H%M%S")}.xlsx'
        return self._finalize(ws, headers, border, font, fill, alignment, filename)

    def _export_stock_in(self, stock_ins):
        wb = Workbook()
        ws, font, fill, alignment, border = self._styles(wb)
        ws.title = '入库记录'
        headers = ['案件编号', '货物编码', '货物名称', '入库数量', '单位',
                   '批次号', '供应商', '操作人', '入库时间']
        ws.append(headers)
        for record in stock_ins:
            goods = record.goods
            unit_name = (goods.variety.category.unit.name
                         if goods and goods.variety and goods.variety.category
                         and goods.variety.category.unit else '')
            ws.append([
                goods.case.case_no if goods and goods.case else '',
                goods.code if goods else '',
                goods.name if goods else '',
                float(record.quantity),
                unit_name,
                record.batch_no,
                record.supplier,
                record.operator.username if record.operator else '',
                record.stock_in_time.strftime('%Y-%m-%d %H:%M:%S')
            ])
        filename = f'入库记录_{datetime.now().strftime("%Y%m%d%H%M%S")}.xlsx'
        return self._finalize(ws, headers, border, font, fill, alignment, filename)

    def _export_stock_out(self, stock_outs):
        wb = Workbook()
        ws, font, fill, alignment, border = self._styles(wb)
        ws.title = '出库记录'
        headers = ['案件编号', '货物编码', '货物名称', '出库数量', '单位',
                   '领用人', '领用部门', '状态', '操作人', '出库时间']
        ws.append(headers)
        for record in stock_outs:
            goods = record.goods
            unit_name = (goods.variety.category.unit.name
                         if goods and goods.variety and goods.variety.category
                         and goods.variety.category.unit else '')
            ws.append([
                goods.case.case_no if goods and goods.case else '',
                goods.code if goods else '',
                goods.name if goods else '',
                float(record.quantity),
                unit_name,
                record.receiver,
                record.receiver_dept,
                record.get_status_display(),
                record.operator.username if record.operator else '',
                record.stock_out_time.strftime('%Y-%m-%d %H:%M:%S')
                if record.stock_out_time else ''
            ])
        filename = f'出库记录_{datetime.now().strftime("%Y%m%d%H%M%S")}.xlsx'
        return self._finalize(ws, headers, border, font, fill, alignment, filename)


class SystemMonitorView(APIView):
    """返回容器内可直接读取的轻量运行状态。"""
    permission_classes = [IsAuthenticated]
    
    def get(self, request):
        cpu_count = os.cpu_count() or 1
        load_1m = os.getloadavg()[0] if hasattr(os, "getloadavg") else 0.0
        disk_total, disk_used, _ = shutil.disk_usage('/')
        disk_percent = round(disk_used * 100 / disk_total, 2) if disk_total else 0
        process_count = len([name for name in os.listdir('/proc') if name.isdigit()]) if os.path.isdir('/proc') else 1
        
        return success_response(data={
            'cpu': {
                'percent': round(load_1m * 100 / cpu_count, 2),
                'count': cpu_count
            },
            'memory': {
                'total': None,
                'used': None,
                'percent': None
            },
            'disk': {
                'total': round(disk_total / (1024 ** 3), 2),
                'used': round(disk_used / (1024 ** 3), 2),
                'percent': disk_percent
            },
            'network': {
                'bytes_sent': None,
                'bytes_recv': None
            },
            'system': {
                'boot_time': None,
                'uptime': None,
                'process_count': process_count
            }
        })
